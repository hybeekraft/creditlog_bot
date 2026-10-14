"""
Script to generate executive Excel spreadsheet of all products from Supplier API and Bot catalog.
"""
import asyncio
import csv
import os
import re
import sqlite3
import subprocess
import sys
from dotenv import load_dotenv

load_dotenv()
from reseller import SupplierRegistry

HAS_OPENPYXL = False
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ModuleNotFoundError:
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "--no-cache-dir", "openpyxl>=3.1.2"], check=True, capture_output=True, timeout=60)
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        HAS_OPENPYXL = True
    except Exception:
        HAS_OPENPYXL = False


def determine_brand_service(name: str) -> str:
    n = name.lower()
    if any(k in n for k in ["chatgpt", "gpt", "openai", "sora"]):
        return "OpenAI / ChatGPT"
    if any(k in n for k in ["claude", "anthropic"]):
        return "Anthropic / Claude"
    if "cursor" in n:
        return "Cursor AI"
    if any(k in n for k in ["gemini", "google"]):
        return "Google / Gemini"
    if "midjourney" in n:
        return "Midjourney"
    if any(k in n for k in ["grok", "xai"]):
        return "xAI / Grok"
    if "deepseek" in n:
        return "DeepSeek"
    if "perplexity" in n:
        return "Perplexity AI"
    if "canva" in n:
        return "Canva"
    if "adobe" in n:
        return "Adobe Creative Cloud"
    if any(k in n for k in ["capcut", "bytedance"]):
        return "CapCut Pro"
    if "telegram" in n:
        return "Telegram Premium"
    if any(k in n for k in ["nordvpn", "surfshark", "expressvpn", "vpn", "proxy"]):
        return "VPN & Proxy Services"
    if any(k in n for k in ["spotify", "netflix", "youtube", "apple music", "disney", "prime"]):
        return "Media & Streaming"
    if any(k in n for k in ["microsoft", "office", "windows", "azure", "onedrive"]):
        return "Microsoft Office / 365"
    if "autodesk" in n:
        return "Autodesk"
    if any(k in n for k in ["github", "copilot"]):
        return "GitHub Copilot"
    if "codex" in n:
        return "Codex API"
    if "higgsfield" in n:
        return "Higgsfield AI"
    if "lovable" in n:
        return "Lovable AI"
    if "elevenlabs" in n:
        return "ElevenLabs AI Voice"
    if any(k in n for k in ["kimi", "moonshot", "glm"]):
        return "Moonshot / Kimi / GLM"
    if "suno" in n:
        return "Suno AI Music"
    if "luma" in n:
        return "Luma Dream Machine"
    if "runway" in n:
        return "Runway Gen-3"
    if "muse" in n:
        return "Muse AI"
    return "AI & Cloud Tools"


def determine_category(name: str, db_cat: str = "") -> str:
    if db_cat and db_cat != "AI Tools":
        return db_cat
    n = name.lower()
    if "cursor" in n:
        return "Cursor API"
    if "codex" in n:
        return "Codex API"
    if "chatgpt" in n:
        return "ChatGPT"
    if "claude" in n:
        return "Claude"
    if any(k in n for k in ["gemini", "google"]):
        return "Gemini"
    if any(k in n for k in ["vpn", "proxy"]):
        return "VPN & Security"
    if any(k in n for k in ["canva", "adobe", "capcut", "midjourney"]):
        return "Creative & Design"
    if any(k in n for k in ["office", "microsoft", "windows"]):
        return "Productivity & Office"
    if any(k in n for k in ["spotify", "netflix", "youtube"]):
        return "Streaming & Entertainment"
    return db_cat or "AI Tools"


def parse_product_desc(name: str, desc: str):
    desc = desc or ""
    lines = [l.strip() for l in desc.splitlines() if l.strip()]

    # 1. Delivery Format
    m_deliv = re.search(r"Delivery format:?\s*([^\n\r]+)", desc, re.IGNORECASE)
    if m_deliv:
        deliv_fmt = m_deliv.group(1).strip()
    elif any(k in desc.lower() for k in ["email | password", "email | pass", "2fa"]):
        deliv_fmt = "Email | Password | 2FA"
    elif "api key" in name.lower() or "tokens" in name.lower() or "credits" in name.lower():
        deliv_fmt = "API Key / Token Secret"
    elif "invite" in name.lower() or "family" in name.lower():
        deliv_fmt = "Activation / Invitation Link"
    elif "key" in name.lower() or "license" in name.lower() or "code" in name.lower():
        deliv_fmt = "License Key / Redemption Code"
    else:
        deliv_fmt = "Account Credentials (Email | Password)"

    # 2. Warranty & Policy
    m_war = re.search(r"((?:No warranty|Warranty:?|Replacement:?)[^\n\r]+)", desc, re.IGNORECASE)
    if m_war:
        warranty = m_war.group(1).strip()
    elif "no warranty" in desc.lower():
        warranty = "No warranty. Delivered as-is; guaranteed active upon initial receipt."
    elif any(k in name.lower() for k in ["1 month", "1m", "30d"]):
        warranty = "24-hour verification & replacement warranty; valid for 30-day billing cycle."
    elif any(k in name.lower() for k in ["1 year", "12m", "12 months"]):
        warranty = "30-day supplier replacement warranty; 12-month subscription period."
    else:
        warranty = "24-hour verification & replacement warranty upon delivery."

    # 3. Critical Handling Notes & Precautions
    m_prec = re.search(r"((?:IMPORTANT|Important information|Notice|CAUTION|Note):?[\s\S]*?)(?=(?:\n[A-Z\s]{4,}:|\Z))", desc, re.IGNORECASE)
    if m_prec:
        precautions = m_prec.group(1).strip()
    elif "chatgpt" in name.lower():
        precautions = (
            "IMPORTANT HANDLING NOTES:\n"
            "1. First, navigate to Settings -> Account -> Email and change the linked email to your personal address.\n"
            "2. Only after updating the email address should you change the password or 2FA authenticator.\n"
            "3. Do not use blacklisted or free public VPN IPs during initial login to prevent automated security flags."
        )
    elif "api" in name.lower():
        precautions = (
            "CRITICAL API SECURITY:\n"
            "1. Never commit API keys to public repositories or client-side code.\n"
            "2. Observe rate limits and token quotas to prevent temporary rate limiting.\n"
            "3. Rotate credentials immediately if any unauthorized exposure is suspected."
        )
    else:
        precautions = (
            "CRITICAL PRECAUTIONS:\n"
            "1. Test and verify login access within 24 hours of delivery.\n"
            "2. Enable personal two-factor authentication (2FA) where supported.\n"
            "3. Avoid simultaneous logins from multiple geographic locations."
        )

    # 4. Instructions
    m_inst = re.search(r"((?:LOGIN INSTRUCTION|HOW TO LOG IN|HOW TO USE|INSTRUCTION|SETUP GUIDE)[\s\S]*?)(?=(?:\nIMPORTANT|\nNOTE|\nWAR|\Z))", desc, re.IGNORECASE)
    if m_inst:
        instructions = m_inst.group(1).strip()
    elif "chatgpt" in name.lower():
        instructions = (
            "STEP-BY-STEP LOGIN INSTRUCTIONS:\n"
            "1. Open the official ChatGPT web portal (https://chatgpt.com) or the official mobile app.\n"
            "2. Enter the delivered email and password.\n"
            "3. When prompted for 2FA, generate a 6-digit code using Google Authenticator or https://2fa.online with the delivered 2FA secret.\n"
            "4. Enter the 6-digit code into ChatGPT (refreshes every 30 seconds).\n"
            "5. Once logged in, bind your personal email in Account Settings."
        )
    elif "api" in name.lower():
        instructions = (
            "API INTEGRATION INSTRUCTIONS:\n"
            "1. Copy the delivered API key / endpoint credentials.\n"
            "2. Set the key in your application environment variable (e.g., OPENAI_API_KEY or BASE_URL).\n"
            "3. Point your HTTP client to the designated endpoint and verify connection with a test ping."
        )
    else:
        instructions = (
            "ACTIVATION INSTRUCTIONS:\n"
            "1. Access the official service portal.\n"
            "2. Log in using the provided credentials or apply the delivered redemption token.\n"
            "3. Confirm premium status activation in your account billing overview."
        )

    # 5. Product Overview & Features
    m_desc = re.search(r"Description:?\s*([^\n\r]+(?:\n[^\n\r]+)?)", desc, re.IGNORECASE)
    if m_desc:
        overview = m_desc.group(1).strip()
    else:
        clean_lines = [l for l in lines if not l.lower().startswith(("item:", "delivery format", "payment method", "no warranty", "login instruction", "how to"))]
        overview = clean_lines[0] if clean_lines else f"Premium digital access for {name}."

    return deliv_fmt, instructions, overview, warranty, precautions


async def build_catalog_excel():
    # 1. Connect to DB to map existing bot IDs and categories
    conn = sqlite3.connect("store.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    
    db_by_sup_id = {}
    db_by_name = {}
    for r in c.execute("SELECT * FROM products").fetchall():
        rd = dict(r)
        if rd.get("supplier_product_id"):
            db_by_sup_id[str(rd["supplier_product_id"])] = rd
        db_by_name[rd["name"].lower().strip()] = rd

    # 2. Fetch live supplier catalog
    registry = SupplierRegistry.from_env()
    client = registry.get()
    supplier_prods = await client.get_products()
    await registry.close_all()
    conn.close()

    if not HAS_OPENPYXL:
        out_file = "CreditLog_Bot_Supplier_Products_Catalog.csv"
        with open(out_file, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([c[0] for c in columns])
            for sp in supplier_prods:
                sid = str(sp["id"])
                sname = sp["name"]
                raw = sp.get("raw", {})
                db_match = db_by_sup_id.get(sid) or db_by_name.get(sname.lower().strip())
                bot_id = db_match["id"] if db_match else f"Auto-{sid}"
                cat_name = determine_category(sname, db_match["category"] if db_match else "")
                brand_svc = determine_brand_service(sname)
                cost_usd = sp["price_usd"]
                store_price = (db_match["price"] / 100.0) if (db_match and db_match.get("price")) else round(cost_usd * 1.45, 2)
                margin_usd = round(store_price - cost_usd, 2)
                stock_val = sp.get("stock")
                if stock_val is not None and stock_val > 0:
                    stock_str = f"In Stock ({stock_val})"
                elif stock_val == 0 or raw.get("inStock") is False:
                    stock_str = "Out of Stock"
                else:
                    stock_str = "Available on Order"
                deliv_fmt, inst, overview, war, prec = parse_product_desc(sname, sp.get("description", ""))
                writer.writerow([bot_id, sname, cat_name, brand_svc, f"${store_price:.2f}", f"${cost_usd:.2f}", f"${margin_usd:.2f}", "Auto-Fulfillment (API)", sid, stock_str, deliv_fmt, inst, overview, war, prec])
        print(f"Catalog fallback CSV saved to {out_file} (openpyxl not installed).")
        return out_file

    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # Styles
    # ----------------------------------------------------
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Dark slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    in_stock_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Light green
    in_stock_font = Font(name="Calibri", size=10, bold=True, color="166534")

    out_stock_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Light red
    out_stock_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    avail_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber
    avail_font = Font(name="Calibri", size=10, bold=True, color="92400E")

    currency_font = Font(name="Calibri", size=10, bold=True, color="047857")
    margin_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    columns = [
        ("Bot ID", 10, "center"),
        ("Product Name", 34, "left"),
        ("Category", 20, "left"),
        ("Brand / Service", 24, "left"),
        ("Store Price (USD)", 16, "right"),
        ("Supplier Cost (USD)", 16, "right"),
        ("Profit Margin ($)", 16, "right"),
        ("Fulfillment Type", 22, "center"),
        ("Supplier ID", 14, "center"),
        ("Stock Status", 18, "center"),
        ("Delivery Format", 26, "left"),
        ("Full Step-by-Step Instructions", 50, "left"),
        ("Product Overview & Features", 40, "left"),
        ("Warranty & Policy", 34, "left"),
        ("Critical Handling Notes & Precautions", 48, "left"),
    ]

    # ----------------------------------------------------
    # Sheet 1: Supplier API Catalog
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "Supplier Products Catalog"
    ws1.views.sheetView[0].showGridLines = True

    # Write Headers
    for col_idx, (col_name, col_width, col_align) in enumerate(columns, 1):
        cell = ws1.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = col_width

    ws1.row_dimensions[1].height = 28

    # Process and write supplier products
    for row_idx, sp in enumerate(supplier_prods, 2):
        sid = str(sp["id"])
        sname = sp["name"]
        raw = sp.get("raw", {})
        
        # Match with DB record
        db_match = db_by_sup_id.get(sid) or db_by_name.get(sname.lower().strip())
        bot_id = db_match["id"] if db_match else f"Auto-{sid}"
        cat_name = determine_category(sname, db_match["category"] if db_match else "")
        brand_svc = determine_brand_service(sname)

        # Price calculations
        cost_usd = sp["price_usd"]
        if db_match and db_match.get("price"):
            store_price = db_match["price"] / 100.0
        else:
            store_price = round(cost_usd * 1.45, 2)

        margin_usd = round(store_price - cost_usd, 2)

        # Stock status
        stock_val = sp.get("stock")
        if stock_val is not None and stock_val > 0:
            stock_str = f"In Stock ({stock_val})"
            stock_style = (in_stock_fill, in_stock_font)
        elif stock_val == 0 or raw.get("inStock") is False:
            stock_str = "Out of Stock"
            stock_style = (out_stock_fill, out_stock_font)
        else:
            stock_str = "Available on Order"
            stock_style = (avail_fill, avail_font)

        deliv_fmt, inst, overview, war, prec = parse_product_desc(sname, sp.get("description", ""))

        row_fill = zebra_fill if (row_idx % 2 == 0) else white_fill

        values = [
            bot_id,
            sname,
            cat_name,
            brand_svc,
            store_price,
            cost_usd,
            margin_usd,
            "Auto-Fulfillment (API)",
            sid,
            stock_str,
            deliv_fmt,
            inst,
            overview,
            war,
            prec,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.fill = row_fill
            cell.font = regular_font

            align_mode = columns[col_idx - 1][2]
            wrap = col_idx in [11, 12, 13, 14, 15]
            cell.alignment = Alignment(horizontal=align_mode, vertical="top", wrap_text=wrap)

            # Currency Formatting
            if col_idx in [5, 6]:
                cell.number_format = "$#,##0.00"
                cell.font = bold_font
            elif col_idx == 7:
                cell.number_format = "$#,##0.00"
                cell.font = currency_font
                cell.fill = margin_fill
            elif col_idx == 10:  # Stock status badge
                cell.fill = stock_style[0]
                cell.font = stock_style[1]

        ws1.row_dimensions[row_idx].height = 42

    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:O{len(supplier_prods)+1}"

    # ----------------------------------------------------
    # Sheet 2: All Bot Products (Complete Store Catalog)
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="All Store Products (DB)")
    ws2.views.sheetView[0].showGridLines = True

    for col_idx, (col_name, col_width, col_align) in enumerate(columns, 1):
        cell = ws2.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = col_width

    ws2.row_dimensions[1].height = 28

    conn = sqlite3.connect("store.db")
    conn.row_factory = sqlite3.Row
    all_db_rows = conn.cursor().execute("SELECT * FROM products ORDER BY id ASC").fetchall()
    conn.close()

    for row_idx, r in enumerate(all_db_rows, 2):
        rd = dict(r)
        sname = rd["name"]
        cat_name = rd.get("category") or determine_category(sname)
        brand_svc = determine_brand_service(sname)
        store_price = (rd.get("price") or 0) / 100.0
        cost_usd = ((rd.get("supplier_price") or 0) / 100.0) if rd.get("supplier_price") else round(store_price * 0.65, 2)
        margin_usd = round(store_price - cost_usd, 2)

        sup_id = str(rd["supplier_product_id"]) if rd.get("supplier_product_id") else "Manual/Local"
        deliv_type = "Auto-Fulfillment (API)" if (rd.get("auto_fulfill") and rd.get("supplier_product_id")) else ("Local Stock" if rd.get("delivery_type") == "auto" else "Manual Delivery")

        # Stock check
        if rd.get("supplier_stock") is not None:
            s_val = rd["supplier_stock"]
            stock_str = f"In Stock ({s_val})" if s_val > 0 else "Out of Stock"
            stock_style = (in_stock_fill, in_stock_font) if s_val > 0 else (out_stock_fill, out_stock_font)
        elif rd.get("supplier_product_id"):
            stock_str = "Available on Order"
            stock_style = (avail_fill, avail_font)
        else:
            stock_str = "Active Store Item"
            stock_style = (in_stock_fill, in_stock_font)

        deliv_fmt, inst, overview, war, prec = parse_product_desc(sname, rd.get("description", ""))

        row_fill = zebra_fill if (row_idx % 2 == 0) else white_fill

        values = [
            r["id"],
            sname,
            cat_name,
            brand_svc,
            store_price,
            cost_usd,
            margin_usd,
            deliv_type,
            sup_id,
            stock_str,
            deliv_fmt,
            inst,
            overview,
            war,
            prec,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.fill = row_fill
            cell.font = regular_font

            align_mode = columns[col_idx - 1][2]
            wrap = col_idx in [11, 12, 13, 14, 15]
            cell.alignment = Alignment(horizontal=align_mode, vertical="top", wrap_text=wrap)

            if col_idx in [5, 6]:
                cell.number_format = "$#,##0.00"
                cell.font = bold_font
            elif col_idx == 7:
                cell.number_format = "$#,##0.00"
                cell.font = currency_font
                cell.fill = margin_fill
            elif col_idx == 10:
                cell.fill = stock_style[0]
                cell.font = stock_style[1]

        ws2.row_dimensions[row_idx].height = 42

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:O{len(all_db_rows)+1}"

    out_file = "CreditLog_Bot_Supplier_Products_Catalog.xlsx"
    wb.save(out_file)
    print(f"Excel catalog successfully saved to {out_file} with {len(supplier_prods)} supplier items and {len(all_db_rows)} total store items.")
    return out_file


if __name__ == "__main__":
    asyncio.run(build_catalog_excel())
